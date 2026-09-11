# -*- coding: utf-8 -*-
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

src = r'C:/Users/kexin.liu/Downloads/CBD_a91766e0077c4994a3f5aa32a0ccdd7e_1788498712908.txt'
dst = r'D:/ai/电商语音ai/CBD_a91766e0077c4994a3f5aa32a0ccdd7e_1788498712908.xlsx'

raw = open(src, 'r', encoding='utf-8-sig').read()
data = json.loads(raw)

wb = Workbook()

# 1) 概览 sheet
ws = wb.active
ws.title = '概览'
ws.append(['项', '值'])
payload = None
for item in data:
    if isinstance(item, dict) and item.get('name') == 'uuid':
        ws.append(['uuid', item.get('value', '')])
    elif isinstance(item, dict) and item.get('name') == 'input_data':
        payload = json.loads(item['value']) if isinstance(item['value'], str) else item['value']

if payload is not None:
    ws.append(['totalTransactionsCount', payload.get('totalTransactionsCount', '')])
    res = payload.get('result') or {}
    ws.append(['resultCode', res.get('resultCode', '')])
    ws.append(['resultTitle', res.get('resultTitle', '')])
    ws.append(['resultDes', res.get('resultDes', '')])

ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 60
for cell in ws['A']:
    cell.font = Font(bold=True)

# 2) 交易明细 sheet
transactions = payload['transactionDetails']
header = list(transactions[0].keys())

ws2 = wb.create_sheet('交易明细')
ws2.append(header)
for t in transactions:
    ws2.append([t.get(k) for k in header])

header_fill = PatternFill('solid', fgColor='4472C4')
for c in ws2[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = header_fill
    c.alignment = Alignment(horizontal='center')
ws2.freeze_panes = 'A2'
for i, k in enumerate(header, 1):
    maxlen = len(str(k))
    for row in ws2.iter_rows(min_row=2, max_row=min(ws2.max_row, 150), min_col=i, max_col=i):
        try:
            v = str(row[0].value or '')
        except Exception:
            v = ''
        maxlen = max(maxlen, len(v))
    ws2.column_dimensions[get_column_letter(i)].width = min(max(maxlen + 2, 10), 40)

wb.save(dst)
print('OK')
print('rows:', len(transactions))
print('cols:', len(header))
print('size:', os.path.getsize(dst))
