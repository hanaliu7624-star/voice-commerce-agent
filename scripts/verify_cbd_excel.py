# -*- coding: utf-8 -*-
from openpyxl import load_workbook
p = r'D:/ai/电商语音ai/CBD_a91766e0077c4994a3f5aa32a0ccdd7e_1788498712908.xlsx'
wb = load_workbook(p, read_only=True)
print('sheets:', wb.sheetnames)
ws = wb['交易明细']
print('rows:', ws.max_row, 'cols:', ws.max_column)
gen = ws.iter_rows(min_row=1, max_row=3, values_only=True)
print('header:', list(next(gen)))
print('first data:', list(next(gen)))
print('second data:', list(next(gen)))
